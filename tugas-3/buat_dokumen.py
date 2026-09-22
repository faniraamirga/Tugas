"""Susun Tugas-3 (Word): rantai nilai dompo pisang berlin Dompis Berlin Thary, Palopo. Hanya yang diminta panduan."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HARGA_AKHIR = 20000
GAMBAR = "tugas-3/rantai_nilai_dompo_pisang_berlin.png"

def rp(n):
    return "Rp" + f"{n:,.0f}".replace(",", ".")

def pct(n, total=HARGA_AKHIR):
    v = n / total * 100
    return (f"{v:.1f}".replace(".", ",") if v != int(v) else f"{int(v)}") + "%"

def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)

def compact(cell):
    for p in cell.paragraphs:
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(0)

def set_grid(tbl, widths):
    tbl.autofit = False
    for row in tbl.rows:
        for c, w in zip(row.cells, widths):
            c.width = w; compact(c)
    # LibreOffice membaca lebar dari tblGrid, bukan dari tiap sel
    for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.twips)))
    for i, row in enumerate(tbl.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if i < len(tbl.rows) - 1:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.5)
    s.left_margin = s.right_margin = Cm(2.5)
st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.line_spacing = 1.5

def para(text="", bold=False, italic=False, align="justify", size=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify": p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    return p

def heading(text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(13)
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    return p

def cell_text(cell, text, bold=False, italic=False, size=10.5, center=False):
    cell.text = ""
    r = cell.paragraphs[0].add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if center: cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

# ---------- Kepala ----------
para("Tugas 3", bold=True, align="center", size=14, space_after=0)
para("Diagram rantai nilai satu produk kreatif lokal", bold=True, align="center", size=13, space_after=0)
para("Dompo pisang berlin Dompis Berlin Thary, Kota Palopo", align="center", size=12, space_after=12)

identitas = [("Nama", "Fanira Amirga"), ("NIM", "……………………"),
             ("Mata kuliah", "Ekonomi Kreatif (MBS12523), pertemuan 3: Ekosistem dan Rantai Nilai"),
             ("Program studi", "Manajemen Bisnis Syariah, FEBI UIN Palopo, semester V"),
             ("Dosen", "Qisty Amalina Rusmana Putri, M.E.")]
t = doc.add_table(rows=len(identitas), cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
for row, (k, v) in zip(t.rows, identitas):
    row.cells[0].text = k; row.cells[1].text = ": " + v
set_grid(t, [Cm(3.5), Cm(12.5)])
para("", space_after=4)

# ---------- Produk ----------
heading("Produk yang dipilih")
para("Dompo pisang berlin merek Dompis Berlin Thary, camilan pisang yang dijemur sampai kering lalu digoreng, "
     "kemasan 200 gram. Usaha ini dijalankan Eka Indah Lestari (Thary) di Jalan Dr. Ratulangi, Kelurahan "
     "Salobulo, Kecamatan Wara Utara, Kota Palopo, dan menjadi UMKM binaan Dinas Koperasi dan UMKM Kota Palopo. "
     "Pembuat dompo lain di Palopo memakai pisang kepok atau raja; Thary memakai pisang berlin yang kecil, manis, "
     "dan sedikit asam, dan memulainya dari pisang lewat matang yang tidak laku di lapak neneknya di pasar "
     "(Nandini, 2024).")

# ---------- Diagram ----------
heading("Rantai nilai dari bahan mentah sampai ke tangan pembeli")
para("Satuan hitungnya satu sisir pisang berlin, kira-kira 1 kilogram, yang setelah dikupas, dijemur dua sampai "
     "tiga hari, dan digoreng menyusut menjadi sekitar 200 gram dompo. Jadi satu sisir pisang menjadi satu "
     "bungkus dompo. Nilai tambah dihitung dengan rumus di kelas: harga jual dikurangi bahan yang dibeli.")
doc.add_picture(GAMBAR, width=Cm(16))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.paragraphs[-1].paragraph_format.keep_with_next = True
para("Gambar 1. Rantai nilai satu bungkus dompo pisang berlin 200 gram (angka perkiraan)",
     italic=True, align="center", size=10, space_after=10)

# ---------- Tabel ----------
baris = [
    ("Petani pisang berlin", "tidak ada", "Rp3.500\n(1 sisir)", 0, 3500),
    ("Pedagang pisang di pasar Palopo", "Rp3.500", "Rp7.500\n(1 sisir)", 3500, 7500),
    ("Dompis Berlin Thary", "Rp7.500\n(1 sisir pisang)", "Rp14.000\n(1 bungkus 200 g)", 7500, 14000),
    ("Toko oleh-oleh Palopo", "Rp14.000", "Rp20.000", 14000, HARGA_AKHIR),
]
tbl = doc.add_table(rows=1, cols=5); tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for c, h in zip(tbl.rows[0].cells, ["Pelaku", "Bahan dibeli", "Dijual seharga", "Nilai tambah", "Bagian dari harga"]):
    cell_text(c, h, bold=True, center=True); shade(c, "D9D9D9")
total_nt = 0
for pelaku, beli_txt, jual_txt, beli, jual in baris:
    nt = jual - beli; total_nt += nt
    cells = tbl.add_row().cells
    for i, v in enumerate([pelaku, beli_txt, jual_txt, rp(nt), pct(nt)]):
        cell_text(cells[i], v, bold=(i == 3), center=(i > 0), size=10)
cells = tbl.add_row().cells
for i, v in enumerate(["Harga di tangan pembeli", "", rp(HARGA_AKHIR), rp(total_nt), "100%"]):
    cell_text(cells[i], v, bold=True, center=(i > 0), size=10); shade(cells[i], "F2F2F2")
set_grid(tbl, [Cm(4.2), Cm(3.4), Cm(3.4), Cm(2.5), Cm(2.5)])
assert total_nt == HARGA_AKHIR
para("Tabel 1. Perkiraan harga dan nilai tambah tiap pelaku per bungkus dompo 200 gram",
     italic=True, align="center", size=10, space_after=8)
para("Angka di atas perkiraan dari harga pisang berlin per tandan di tingkat petani dan pedagang (Hairon, 2022; "
     "Jatimnow, 2020), harga pisang per sisir di pasar Makassar 2026 (Macassar.id, 2026), harga dompo per kepak di "
     "pasar dan setelah dikemas (ZonaSultra, 2019), harga kemasan standing pouch (Kemasan Retail, 2021), serta "
     "komposisi biaya sale pisang skala rumah tangga (Yandra dkk., 2025). Angka sebenarnya akan saya tanyakan "
     "langsung ke Thary, pedagang pisang, dan toko oleh-oleh.", italic=True, size=10)

# ---------- Paragraf penutup ----------
heading("Siapa yang bagiannya paling kecil, dan apakah itu sepadan?")
para("Bagian paling kecil diambil petani, Rp3.500 dari Rp20.000, untuk menanam dan menunggu hampir setahun "
     "sampai satu tandan siap panen, lalu menjualnya dalam hitungan hari dengan harga berapa pun yang ditawarkan "
     "pedagang karena pisang matang tidak bisa disimpan. Menurut saya bagian itu terlalu kecil untuk waktu yang "
     "ditanggung, walaupun biaya tunai petani hampir nol dan pisang berlin hanya tanaman sela di kebunnya. Angka "
     "yang menurut saya paling tidak sepadan ada di baris Thary. Nilai tambahnya tampak paling besar, "
     "Rp6.500, tetapi dari angka itu sekitar Rp3.300 habis untuk tepung, minyak, gas, kemasan, dan stiker, "
     "sehingga yang tersisa untuk kerja mengupas belasan buah, menjemur dua sampai tiga hari sambil menjaga dari "
     "hujan, menggoreng, mengemas, dan mengurus izin PIRT hanya sekitar Rp3.200. Toko oleh-oleh yang hanya "
     "memajang mendapat Rp6.000 tanpa biaya bahan apa pun, dan pedagang pasar yang memegang pisang dua hari "
     "mendapat Rp4.000. Toko memang membayar sewa dan menanggung barang yang tidak laku, tetapi bagiannya besar "
     "karena ia memegang informasi pasar: ia tahu siapa perantau yang mencari oleh-oleh dan berapa mereka sanggup "
     "bayar, sementara petani di kebun dan Thary di dapur tidak. Ini pola yang sama dengan bakul rotan di kelas, "
     "semakin dekat ke konsumen semakin besar nilai yang diambil. Thary sudah menemukan jalan keluarnya: ketika ia "
     "menjual sendiri lewat WhatsApp dan Instagram dengan harga yang sama, bagian toko Rp6.000 berpindah ke "
     "tangannya tanpa pembeli membayar lebih. Petani belum punya jalan seperti itu, dan di sanalah bagian yang "
     "paling perlu dinaikkan.")

# ---------- Referensi ----------
heading("Referensi")
refs = [
    "Hairon, I. (2022, 22 Juni). Di Jember, pisang berlin dijual sangat murah. Suara Indonesia. "
    "https://suaraindonesia.co.id/news/features/62b2aebc1a9e4/di-jember-pisang-berlin-dijual-sangat-murah",
    "Jatimnow. (2020, 29 Oktober). Imut dan rasa sedikit masam, pisang barlin di PIOS diburu pengunjung. "
    "https://jatimnow.com/baca-30937-imut-dan-rasa-sedikit-masam-pisang-barlin-di-pios-diburu-pengunjung",
    "Kemasan Retail. (2021). Kemasan standing pouch kraft window zipper [Katalog produk]. "
    "https://kemasanretail.com/product/kemasan-standing-pouch-kraft-window-zipper-16x24-750g/",
    "Macassar.id. (2026, 20 Februari). Ini harga bahan pangan saat Ramadhan. "
    "https://macassar.id/2026/02/20/ini-harga-bahan-pangan-saat-ramadhan/",
    "Nandini, A. B. (2024, 12 November). Dompis Berlin Thary, olahan pisang khas Kota Palopo Sulsel yang bikin "
    "nagih. Tribun Timur. https://makassar.tribunnews.com/2024/11/12/dompis-berlin-thary-olahan-pisang-khas-kota-"
    "palopo-sulsel-yang-bikin-nagih",
    "Yandra, H. M., Hakimi, R., & Raesi, S. (2025). Analisis nilai tambah pengolahan pisang (Musa paradisiaca L.) "
    "pada UMKM Heni Pisang Sale di Nagari IV Koto, Kecamatan Pulau Punjung, Kabupaten Dharmasraya. JAS (Jurnal "
    "Agri Sains), 9(1). https://ojs.umb-bungo.ac.id/index.php/JAS/article/download/1725/1384",
    "ZonaSultra. (2019, 9 Mei). Dompo pisang, camilan khas Bombana yang mulai dikembangkan. "
    "https://zonasultra.id/dompo-pisang-camilan-khas-bombana-yang-mulai-dikembangkan/",
]
for r_ in refs:
    p = para(r_, size=11, space_after=4)
    p.paragraph_format.left_indent = Cm(1); p.paragraph_format.first_line_indent = Cm(-1)
    p.paragraph_format.line_spacing = 1.15

doc.save("tugas-3/Tugas-3_Rantai_Nilai_Dompo_Pisang_Berlin.docx")
print("saved")
